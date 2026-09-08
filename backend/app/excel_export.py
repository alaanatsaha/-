"""Append completed assessments to the bank-branded Excel register."""
from __future__ import annotations

import logging
import threading
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo

from . import models

logger = logging.getLogger(__name__)
_EXPORT_LOCK = threading.Lock()
_EXPORT_PATH = Path(__file__).resolve().parents[1] / "data" / "assessments.xlsx"
_HEADERS = [
    "رقم الدراسة",
    "تاريخ الدراسة",
    "اسم العميل",
    "العمر",
    "رقم الحساب",
    "الجهة",
    "البطاقة المطلوبة",
    "الحد الائتماني",
    "الراتب الشهري",
    "الدخل الإضافي",
    "الالتزامات الحالية",
    "الدخل المعتمد",
    "عبء البطاقة الشهري",
    "إجمالي العبء الشهري",
    "نسبة عبء الدين DPR",
    "الحد الأقصى DPR",
    "تحويل الراتب",
    "كفيل محول راتب",
    "النتيجة",
    "ملخص القرار",
    "ملاحظات الموظف",
]

_BLUE = "003B70"
_BLUE_LIGHT = "EAF3F9"
_GOLD = "F4C400"
_BORDER = Side(style="thin", color="DCE5EC")


def append_assessment(assessment: models.Assessment) -> Path:
    """Append one committed assessment and return the workbook path."""
    with _EXPORT_LOCK:
        _EXPORT_PATH.parent.mkdir(parents=True, exist_ok=True)
        if _EXPORT_PATH.exists():
            workbook = load_workbook(_EXPORT_PATH)
            worksheet = workbook.active
        else:
            workbook = Workbook()
            worksheet = workbook.active
            worksheet.title = "سجل التقييمات"
            worksheet.append(_HEADERS)
            _style_header(worksheet)

        if worksheet.max_row == 1 and worksheet.cell(1, 1).value is None:
            worksheet.delete_rows(1)
            worksheet.append(_HEADERS)
            _style_header(worksheet)

        worksheet.append(_assessment_row(assessment))
        row_number = worksheet.max_row
        _style_data_row(worksheet, row_number)
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        worksheet.sheet_view.rightToLeft = True
        _set_column_widths(worksheet)
        _ensure_table(worksheet)
        workbook.save(_EXPORT_PATH)
        return _EXPORT_PATH


def get_export_path() -> Path:
    return _EXPORT_PATH


def _assessment_row(assessment: models.Assessment) -> list[object]:
    return [
        assessment.id,
        assessment.created_at.strftime("%Y-%m-%d %H:%M"),
        assessment.customer_name,
        assessment.age,
        assessment.account_number,
        assessment.employer.name,
        assessment.card.name,
        assessment.card.credit_limit,
        assessment.salary,
        assessment.other_income,
        assessment.existing_facilities,
        assessment.approved_income,
        assessment.card.monthly_burden,
        assessment.total_monthly_burden,
        assessment.dpr / 100,
        assessment.max_dpr_applied / 100,
        "نعم" if assessment.salary_transferred else "لا",
        "نعم" if assessment.guarantor_available else "لا",
        "مؤهل مبدئيًا" if assessment.status == models.AssessmentStatus.eligible else "غير مؤهل مبدئيًا",
        assessment.decision_summary,
        assessment.notes or "",
    ]


def _style_header(worksheet) -> None:
    for cell in worksheet[1]:
        cell.fill = PatternFill("solid", fgColor=_BLUE)
        cell.font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=Side(style="medium", color=_GOLD))
    worksheet.row_dimensions[1].height = 32


def _style_data_row(worksheet, row_number: int) -> None:
    for cell in worksheet[row_number]:
        cell.font = Font(name="Arial", size=10, color="18324A")
        cell.alignment = Alignment(horizontal="right", vertical="center", wrap_text=True)
        cell.border = Border(bottom=_BORDER)
    for column in (8, 9, 10, 11, 12, 13, 14):
        worksheet.cell(row_number, column).number_format = '#,##0.00'
    for column in (15, 16):
        worksheet.cell(row_number, column).number_format = '0.00%'
    if row_number % 2 == 0:
        for cell in worksheet[row_number]:
            cell.fill = PatternFill("solid", fgColor=_BLUE_LIGHT)


def _set_column_widths(worksheet) -> None:
    widths = [12, 18, 24, 8, 18, 24, 18, 14, 14, 14, 16, 14, 17, 17, 15, 15, 14, 16, 18, 38, 32]
    for index, width in enumerate(widths, start=1):
        worksheet.column_dimensions[worksheet.cell(1, index).column_letter].width = width


def _ensure_table(worksheet) -> None:
    if worksheet.tables:
        next(iter(worksheet.tables.values())).ref = f"A1:U{worksheet.max_row}"
        return
    table = Table(displayName="AssessmentsRegister", ref=f"A1:U{worksheet.max_row}")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    worksheet.add_table(table)
